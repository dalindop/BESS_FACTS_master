# -*- coding: utf-8 -*-
"""
Test de parameters.py usando los datos REALES del IEEE 6-bus
(Wood & Wollenberg) leidos del Excel. Construye un `data` minimo
imitando lo que hara el data_loader, y valida los parametros.
"""

import openpyxl
import pyomo.environ as pyo
import sets
import parameters as par


# ----------------------------------------------------------------------
# Mini-loader SOLO para el test (el loader real ira en 01_Inputs/)
# ----------------------------------------------------------------------
XLSX = "/mnt/user-data/uploads/Case_IEEE6.xlsx"
N_HORAS = 4          # pocas horas para test rapido
N_SEG_COST = 3       # tramos de linealizacion de costo


def leer_hoja(ws):
    """Devuelve (encabezado, lista_de_filas) de una hoja."""
    filas = list(ws.iter_rows(values_only=True))
    return filas[0], filas[1:]


class Data:
    pass


wb = openpyxl.load_workbook(XLSX, data_only=True)
data = Data()
data.mva_base = 100
data.n_horas = N_HORAS
data.n_seg_costo = N_SEG_COST
data.n_seg_perdidas = 3

# --- Buses: ids de texto 'b1','b2',... y demanda fija Pd --------------
hdr, filas = leer_hoja(wb["Buses"])
col = {name: i for i, name in enumerate(hdr)}
data.nodos = []
pd_por_nodo = {}
for f in filas:
    nid = f"b{f[col['bus_i']]}"          # 1 -> 'b1'
    data.nodos.append(nid)
    pd_por_nodo[nid] = f[col["Pd_MW"]]

# --- Lineas: ids 'l1','l2',... ; susceptancia=1/x ; g=r/(r^2+x^2) -----
hdr, filas = leer_hoja(wb["Lineas"])
col = {name: i for i, name in enumerate(hdr)}
data.lineas = []
data.susceptancia, data.conductancia, data.flow_max = {}, {}, {}
data.linea_from, data.linea_to = {}, {}
for k, f in enumerate(filas, start=1):
    lid = f"l{k}"
    r = f[col["r_pu"]]
    x = f[col["x_pu"]]
    data.lineas.append(lid)
    data.susceptancia[lid] = 1.0 / x
    data.conductancia[lid] = r / (r**2 + x**2)
    data.flow_max[lid] = f[col["rateA_MW"]]
    data.linea_from[lid] = f"b{f[col['from_bus']]}"
    data.linea_to[lid]   = f"b{f[col['to_bus']]}"

# --- Generadores: ids 'G1',... ; limites ; rampas --------------------
hdr, filas = leer_hoja(wb["Generadores"])
col = {name: i for i, name in enumerate(hdr)}
data.gen_termica = []
data.pmin_term, data.pmax_term = {}, {}
data.ramp_up, data.ramp_down = {}, {}
gen_orden = []
for k, f in enumerate(filas, start=1):
    gid = f"G{k}"
    data.gen_termica.append(gid)
    gen_orden.append(gid)
    data.pmin_term[gid] = f[col["Pmin_MW"]]
    data.pmax_term[gid] = f[col["Pmax_MW"]]
    # el W&W no trae rampa horaria util; usamos el rango como cota amplia
    data.ramp_up[gid]   = f[col["rango_MW"]]
    data.ramp_down[gid] = f[col["rango_MW"]]

# --- Costos: polinomio c2,c1,c0 -> linealizar por tramos -------------
hdr, filas = leer_hoja(wb["Costos"])
col = {name: i for i, name in enumerate(hdr)}
data.slope_term, data.fg_min_term = {}, {}
for gid, f in zip(gen_orden, filas):
    c2 = f[col["c2_$/MW2"]]
    c1 = f[col["c1_$/MW"]]
    c0 = f[col["c0_$"]]
    fgmin, slopes = par.linealizar_costo(
        c2, c1, c0, data.pmin_term[gid], data.pmax_term[gid], N_SEG_COST)
    data.fg_min_term[gid] = fgmin
    for m, sl in enumerate(slopes, start=1):
        data.slope_term[(gid, m)] = sl

# --- Demanda horaria: constante en el tiempo (fase validacion) -------
data.demanda = {(n, t): pd_por_nodo[n]
                for n in data.nodos for t in range(1, N_HORAS + 1)}

# (no hay hidro/renov/bess/facts en el W&W: quedan vacios)
data.gen_hidro, data.gen_renov = [], []
data.lineas_cand, data.nodos_bess, data.lineas_facts = [], [], []


# ----------------------------------------------------------------------
# Construir modelo y validar
# ----------------------------------------------------------------------
model = pyo.ConcreteModel()
sets.build_sets(model, data)
par.build_parameters(model, data)

assert pyo.value(model.MVA_base) == 100
assert len(model.L) == 11
assert len(model.G) == 3
# susceptancia l1: x=0.2 -> 1/0.2 = 5
assert abs(pyo.value(model.susceptance["l1"]) - 5.0) < 1e-9
# flujo max l1 = rateA = 40
assert pyo.value(model.flow_max["l1"]) == 40
# demanda nodo b4 = 70 en cualquier hora
assert pyo.value(model.D["b4", 1]) == 70
assert pyo.value(model.D["b4", N_HORAS]) == 70
# demanda nodos generadores = 0
assert pyo.value(model.D["b1", 1]) == 0
# fg_min de G1 (costo en Pmin=50) ~ 809.875
assert abs(pyo.value(model.fg_min["G1"]) - 809.875) < 0.01
# pendientes crecientes en G1 (costo convexo)
s1 = pyo.value(model.slope["G1", 1])
s3 = pyo.value(model.slope["G1", 3])
assert s3 > s1, "las pendientes deben crecer (costo convexo)"

print("OK: parametros construidos desde el Excel REAL del W&W y validados.")
print(f"  Nodos   : {data.nodos}")
print(f"  Lineas  : {len(data.lineas)}  (susceptance l1 = {pyo.value(model.susceptance['l1']):.2f})")
print(f"  Gen     : {data.gen_termica}")
print(f"  Demanda total por hora = {sum(pd_por_nodo.values())} MW")
print(f"  fg_min  : " + ", ".join(f"{g}={pyo.value(model.fg_min[g]):.1f}" for g in model.G))
print(f"  slope G1: {[round(pyo.value(model.slope['G1',m]),3) for m in model.SEG_COST]}")


# ----------------------------------------------------------------------
# Bloque 4 -- costos de inversion
# ----------------------------------------------------------------------
# Con el W&W los conjuntos LC y F estan vacios, asi que Cl, Cf quedan
# vacios (sin error). Los costos escalares de BESS toman los defaults
# de Alvaro (45/45).
assert len(model.Cl) == 0, "LC vacio en W&W -> Cl vacio (sin error)"
assert pyo.value(model.Cs_power) == 45,  "costo potencia BESS (Alvaro)"
assert pyo.value(model.Cs_energy) == 45, "costo energia BESS (Alvaro)"
assert pyo.value(model.Cs_inst) == 0,    "costo inst BESS (no en Alvaro)"
assert pyo.value(model.Cf_inst) == 0,    "costo inst FACTS (placeholder)"
assert pyo.value(model.Cf_size) == 0,    "costo size FACTS (placeholder)"

print()
print("OK: Bloque 4 (costos de inversion) construido y validado.")
print(f"  Cl (lineas cand.) : {len(model.Cl)} entradas (LC vacio en W&W)")
print(f"  BESS  : inst={pyo.value(model.Cs_inst)}, "
      f"power={pyo.value(model.Cs_power)}, energy={pyo.value(model.Cs_energy)}")
print(f"  FACTS : inst={pyo.value(model.Cf_inst)}, "
      f"size={pyo.value(model.Cf_size)} (placeholders)")

# --- Prueba extra: que SI toma costos personalizados cuando se proveen
import copy
d2 = copy.deepcopy(data)          # copia la instancia ya poblada
# anadir LC con un costo, y costos FACTS
d2.lineas_cand = ["lc1"]
d2.costo_linea = {"lc1": 1_000_000}
d2.costo_facts_inst = 200_000
m2 = pyo.ConcreteModel()
sets.build_sets(m2, d2)
par.build_parameters(m2, d2)
assert pyo.value(m2.Cl["lc1"]) == 1_000_000
assert pyo.value(m2.Cf_inst) == 200_000
print("OK: costos personalizados (linea candidata + FACTS) se aplican bien.")
